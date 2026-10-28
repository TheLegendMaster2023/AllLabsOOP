import os
import re
import docx
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

REPORTS_DIR = r"d:\AllLabsOOP\Reports"
SCREENS_DIR = os.path.join(REPORTS_DIR, "screens")

LAB_INFO = {
    1: {"topic": "Разработка консольного приложения"},
    2: {"topic": "Перенаправление потоков ввода-вывода"},
    3: {"topic": "Управление потоком выполнения в программе"},
    4: {"topic": "Одномерные массивы"},
    5: {"topic": "Многомерные массивы"},
    6: {"topic": "Классы. Объектное моделирование"},
    7: {"topic": "Конструктор класса. Перегрузка методов класса"},
    8: {"topic": "Проектирование иерархии классов"},
    9: {"topic": "Полиморфизм на основе интерфейсов"},
    10: {"topic": "Основы работы с файлами. Многомодульность"},
    11: {"topic": "Обработка данных в файлах (CSV и LINQ)"},
    12: {"topic": "Событийно-ориентированное программирование в графическом интерфейсе Windows Forms"}
}

def clean_latex(text):
    if not text:
        return ""
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    text = text.replace(r'\cdot', '·')
    text = text.replace(r'\times', '×')
    text = text.replace(r'\approx', '≈')
    text = text.replace(r'\sqrt', '√')
    text = text.replace(r'\sin', 'sin')
    text = text.replace(r'\cos', 'cos')
    text = text.replace(r'\exp', 'exp')
    text = text.replace(r'\le', '≤')
    text = text.replace(r'\ge', '≥')
    text = text.replace(r'\pm', '±')
    text = text.replace(r'\to', '→')
    text = text.replace(r'\rightarrow', '→')
    text = text.replace(r'\left(', '(').replace(r'\right)', ')')
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = text.replace('$$', '').replace('$', '')
    return text

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
                        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

def add_page_number_to_section(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def create_title_page(doc, lab_num, topic):
    # Paragraph 1: Шапка министерства
    p_min = doc.add_paragraph()
    p_min.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_min.paragraph_format.line_spacing = 1.0
    p_min.paragraph_format.space_before = Pt(0)
    p_min.paragraph_format.space_after = Pt(2)
    p_min.paragraph_format.first_line_indent = Mm(0)
    
    run_min = p_min.add_run("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ\n"
                            "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ\n"
                            "ВЫСШЕГО ОБРАЗОВАНИЯ\n"
                            "«СЕВЕРО-КАВКАЗСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»")
    run_min.font.name = 'Times New Roman'
    run_min.font.size = Pt(11)
    run_min.font.bold = True
    run_min.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 2: Институт и кафедра
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_before = Pt(4)
    p_inst.paragraph_format.space_after = Pt(65)
    p_inst.paragraph_format.first_line_indent = Mm(0)
    
    run_inst = p_inst.add_run("Институт цифрового развития\n"
                              "Кафедра прикладной информатики")
    run_inst.font.name = 'Times New Roman'
    run_inst.font.size = Pt(12)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 3: Название отчета
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.first_line_indent = Mm(0)
    
    run_otchet = p_title.add_run("ОТЧЕТ\n")
    run_otchet.font.name = 'Times New Roman'
    run_otchet.font.size = Pt(18)
    run_otchet.font.bold = True
    run_otchet.font.color.rgb = RGBColor(0, 0, 0)
    
    run_lab = p_title.add_run(f"по лабораторной работе №{lab_num}\n"
                              f"по дисциплине: «Объектно-ориентированное программирование»\n")
    run_lab.font.name = 'Times New Roman'
    run_lab.font.size = Pt(14)
    run_lab.font.bold = True
    run_lab.font.color.rgb = RGBColor(0, 0, 0)

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.line_spacing = 1.15
    p_topic.paragraph_format.space_before = Pt(0)
    p_topic.paragraph_format.space_after = Pt(65)
    p_topic.paragraph_format.first_line_indent = Mm(0)
    
    run_topic = p_topic.add_run(f"Тема: «{topic}»\n"
                                f"Вариант №7")
    run_topic.font.name = 'Times New Roman'
    run_topic.font.size = Pt(14)
    run_topic.font.bold = True
    run_topic.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 4: Сведения о студенте
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_author.paragraph_format.left_indent = Mm(90)
    p_author.paragraph_format.line_spacing = 1.15
    p_author.paragraph_format.space_before = Pt(0)
    p_author.paragraph_format.space_after = Pt(90)
    p_author.paragraph_format.first_line_indent = Mm(0)
    
    run_author = p_author.add_run("Выполнил:\n"
                                  "студент группы ПИН-б-о-24-1\n"
                                  "направления подготовки\n"
                                  "09.03.03 «Прикладная информатика»\n"
                                  "Мухамеджанов Э. Т.")
    run_author.font.name = 'Times New Roman'
    run_author.font.size = Pt(13)
    run_author.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 5: Город и год
    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.line_spacing = 1.0
    p_city.paragraph_format.space_before = Pt(0)
    p_city.paragraph_format.space_after = Pt(0)
    p_city.paragraph_format.first_line_indent = Mm(0)
    
    run_city = p_city.add_run("Ставрополь, 2026 г.")
    run_city.font.name = 'Times New Roman'
    run_city.font.size = Pt(14)
    run_city.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_paragraph(doc, text, bold_prefix="", indent=True):
    text = clean_latex(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Mm(12.5) if indent else Mm(0)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(14)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)

    # Process inline formatting: bold **...** and code `...`
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(20, 20, 20)
        else:
            run = p.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_block(doc, code_str, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.first_line_indent = Mm(12.5)
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(3)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run(caption)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(50, 50, 50)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_borders(cell)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    cell.width = Mm(165)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(20, 20, 20)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(0)
    p_after.paragraph_format.line_spacing = 1.0

def add_image_with_caption(doc, img_rel_path, caption="", width_mm=150):
    full_path = os.path.normpath(os.path.join(REPORTS_DIR, img_rel_path))
    if not os.path.exists(full_path):
        full_path = os.path.normpath(os.path.join(SCREENS_DIR, os.path.basename(img_rel_path)))
    
    if os.path.exists(full_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.first_line_indent = Mm(0)
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        
        run = p_img.add_run()
        run.add_picture(full_path, width=Mm(width_mm))
        
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.first_line_indent = Mm(0)
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            r = p_cap.add_run(caption)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(40, 40, 40)

def convert_report(lab_num):
    info = LAB_INFO[lab_num]
    topic = info["topic"]
    md_path = os.path.join(REPORTS_DIR, f"LR{lab_num:02d}_Report.md")
    out_docx = os.path.join(REPORTS_DIR, f"LR{lab_num:02d}_Report.docx")
    
    print(f"Converting Lab {lab_num}: {topic}...")
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.different_first_page_header_footer = True
    
    add_page_number_to_section(section)
    create_title_page(doc, lab_num, topic)
    
    # Strip title page from markdown (skip everything until '## 1.')
    idx_start = md_text.find("## 1.")
    if idx_start == -1:
        idx_start = md_text.find("## 1")
    if idx_start != -1:
        content = md_text[idx_start:]
    else:
        content = md_text

    # Parse content line by line or block by block
    lines = content.split('\n')
    i = 0
    num_lines = len(lines)
    
    in_code_block = False
    code_buf = []
    code_lang = ""
    
    while i < num_lines:
        line = lines[i]
        stripped = line.strip()
        
        # Check code fence
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_buf = []
                i += 1
                continue
            else:
                in_code_block = False
                code_text = "\n".join(code_buf)
                
                # Check if it was a mermaid block
                if code_lang == "mermaid":
                    # Embed UML image
                    if lab_num == 6:
                        add_image_with_caption(doc, "screens/lr6_uml.png", "Рисунок 6.1 – UML-диаграмма класса MatrixMxN", width_mm=140)
                    elif lab_num == 8:
                        add_image_with_caption(doc, "screens/lr8_uml.png", "Рисунок 8.1 – UML-диаграмма иерархии классов Машины", width_mm=155)
                    elif lab_num == 9:
                        add_image_with_caption(doc, "screens/lr9_uml.png", "Рисунок 9.1 – UML-диаграмма интерфейсов и классов", width_mm=150)
                else:
                    add_code_block(doc, code_text)
                i += 1
                continue
                
        if in_code_block:
            code_buf.append(line)
            i += 1
            continue
            
        # Skip empty lines or horizontal rules
        if not stripped or stripped == "---":
            i += 1
            continue
            
        # Check headings
        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            add_heading_1(doc, h_text)
            i += 1
            continue
        elif stripped.startswith("### "):
            h_text = stripped[4:].strip()
            add_heading_2(doc, h_text)
            i += 1
            continue
        elif stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            add_heading_3(doc, h_text)
            i += 1
            continue
            
        # Check table
        if stripped.startswith("|") and stripped.endswith("|"):
            # Collect all table lines
            table_lines = []
            while i < num_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            # Parse table rows
            parsed_rows = []
            for tl in table_lines:
                # check if separator
                if re.match(r'^\|[\s\-:|]+\|$', tl):
                    continue
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                parsed_rows.append(cells)
                
            if parsed_rows:
                # Check if table contains images
                has_imgs = any('![' in cell for row in parsed_rows for cell in row)
                if has_imgs and len(parsed_rows) >= 2:
                    # 2-column image comparison table
                    col_count = len(parsed_rows[0])
                    table = doc.add_table(rows=len(parsed_rows), cols=col_count)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row in enumerate(parsed_rows):
                        for c_idx, cell_text in enumerate(row):
                            cell = table.cell(r_idx, c_idx)
                            set_cell_borders(cell)
                            cell_w_mm = int(160 / col_count)
                            cell.width = Mm(cell_w_mm)
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.0
                            cp.paragraph_format.space_before = Pt(2)
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.first_line_indent = Mm(0)
                            
                            img_match = re.search(r'!\[(.*?)\]\((.*?)\)', cell_text)
                            if img_match:
                                img_cap = img_match.group(1)
                                img_src = img_match.group(2)
                                fpath = os.path.normpath(os.path.join(REPORTS_DIR, img_src))
                                if not os.path.exists(fpath):
                                    fpath = os.path.normpath(os.path.join(SCREENS_DIR, os.path.basename(img_src)))
                                if os.path.exists(fpath):
                                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    r = cp.add_run()
                                    r.add_picture(fpath, width=Mm(cell_w_mm - 6))
                                    if img_cap:
                                        p_c = cell.add_paragraph()
                                        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                        p_c.paragraph_format.first_line_indent = Mm(0)
                                        r2 = p_c.add_run(img_cap)
                                        r2.font.name = 'Times New Roman'
                                        r2.font.size = Pt(10)
                                        r2.font.italic = True
                            else:
                                if r_idx == 0:
                                    set_cell_background(cell, "F0F0F0")
                                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    r = cp.add_run(clean_latex(cell_text))
                                    r.font.name = 'Times New Roman'
                                    r.font.size = Pt(11)
                                    r.font.bold = True
                                else:
                                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                    r = cp.add_run(clean_latex(cell_text))
                                    r.font.name = 'Times New Roman'
                                    r.font.size = Pt(11)
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_before = Pt(4)
                else:
                    # Regular text table
                    col_count = len(parsed_rows[0])
                    table = doc.add_table(rows=len(parsed_rows), cols=col_count)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row in enumerate(parsed_rows):
                        for c_idx, cell_text in enumerate(row):
                            if c_idx >= col_count:
                                continue
                            cell = table.cell(r_idx, c_idx)
                            set_cell_borders(cell)
                            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.15
                            cp.paragraph_format.space_before = Pt(2)
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.first_line_indent = Mm(0)
                            if r_idx == 0:
                                set_cell_background(cell, "F0F0F0")
                                r = cp.add_run(clean_latex(cell_text))
                                r.font.name = 'Times New Roman'
                                r.font.size = Pt(11)
                                r.font.bold = True
                            else:
                                r = cp.add_run(clean_latex(cell_text))
                                r.font.name = 'Times New Roman'
                                r.font.size = Pt(11)
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_before = Pt(4)
            continue

        # Check image standalone
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            img_cap = img_match.group(1)
            img_src = img_match.group(2)
            
            # Look ahead for caption text like *Рисунок 1.1 – ...*
            caption_line = ""
            if i + 1 < num_lines and lines[i+1].strip().startswith("*Рисунок"):
                caption_line = lines[i+1].strip().strip("*")
                i += 1
            elif img_cap:
                caption_line = img_cap
                
            add_image_with_caption(doc, img_src, caption_line)
            i += 1
            continue

        # Check question & answer or list item
        # Question format in section 3: **1. Question?**
        q_match = re.match(r'^\*\*(\d+\..*?)\*\*\s*(.*)$', stripped)
        if q_match:
            q_num = q_match.group(1)
            q_rest = q_match.group(2)
            if q_rest:
                add_body_paragraph(doc, q_rest, bold_prefix=q_num + " ")
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.first_line_indent = Mm(12.5)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(q_num)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)
                r.font.bold = True
            i += 1
            continue

        # Bullet list or numbered list
        list_match = re.match(r'^(\*|-|\d+\.)\s+(.*)$', stripped)
        if list_match:
            bullet = list_match.group(1)
            item_text = list_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Mm(12.5)
            
            r_b = p.add_run(bullet + " ")
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(14)
            r_b.font.bold = True if bullet[0].isdigit() else False
            
            # format item_text
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', clean_latex(item_text))
            for part in parts:
                if not part: continue
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(14)
                    run.font.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(11)
                else:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(14)
            i += 1
            continue

        # Normal paragraph
        add_body_paragraph(doc, stripped)
        i += 1

    doc.save(out_docx)
    print(f"  -> Saved {out_docx}")

if __name__ == "__main__":
    for num in range(1, 13):
        convert_report(num)
    print("\nALL 12 DOCX REPORTS CREATED SUCCESSFULLY!")
